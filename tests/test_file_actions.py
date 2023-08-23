import unittest
from ..actions.file_actions import extract_data_from_excel, generate_word_files, replace_placeholders_in_table
from docx import Document


class TestFileActions(unittest.TestCase):
    def setUp(self):
        # Set up any common data or configurations needed for tests
        self.excel_file_path = "input_data/test_data.xlsx"
        self.template_dir = "note_templates"
        self.output_dir = "output_data"
        
    def test_extract_data_from_excel(self):
        data_list = extract_data_from_excel(self.excel_file_path)
        self.assertIsInstance(data_list, list)
        # Add more specific assertions based on your data structure
        
    def test_generate_word_files(self):
        # Prepare test data
        test_data = [
            ["template_1.docx", "value1", "value2", "value3"],
            ["template_2.docx", "value4", "value5", "value6"]
        ]
        
        # Generate Word files using test data
        generate_word_files(test_data, self.template_dir, self.output_dir)
        
        # Perform assertions to verify the generated files
        
    def test_replace_placeholders_in_table(self):
        # Prepare test data and Word document
        data_row = ["value1", "value2", "value3"]
        doc = Document("path_to_test_document.docx")
        
        # Replace placeholders using the function
        replace_placeholders_in_table(doc, data_row)
        
        # Perform assertions to check if placeholders are replaced correctly
        
if __name__ == "__main__":
    unittest.main()